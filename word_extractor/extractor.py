from docx import Document
import re


class WordExtractor:
    def __init__(self, filepath):
        self.filepath = filepath
        self.doc = Document(filepath)

    def get_title(self):
        if self.doc.core_properties.title:
            return self.doc.core_properties.title

        for para in self.doc.paragraphs[:5]:
            if para.text.strip() and len(para.text.strip()) > 10:
                return para.text.strip()[:200]
        return "Sin título"

    def get_author(self):
        # 1. Look for "Autor (a):" pattern in document FIRST (more reliable)
        for para in self.doc.paragraphs:
            text = para.text.strip()
            if text.startswith("Autor"):
                # Extract name after "Autor (a):" or "Autor:"
                match = re.search(
                    r"(?:Autor\s*(?:\([^)]+\))?:?)\s*(.+?)(?:\n|$)", text, re.IGNORECASE
                )
                if match:
                    return match.group(1).strip()

        # 2. Check core properties (often contains email, not name)
        if self.doc.core_properties.author:
            # If it looks like an email, try to find real name from document
            if "@" not in self.doc.core_properties.author:
                return self.doc.core_properties.author

        # 3. Fallback: look for "presentado por" pattern
        for para in self.doc.paragraphs[:20]:
            text = para.text.lower()
            if "presentado por" in text:
                match = re.search(r"presentado por\s*:?\s*(.+?)(?:\n|$)", text)
                if match:
                    return match.group(1).strip()

        return "Autor desconocido"

    def get_abstract(self):
        text = ""
        capturing = False
        keywords = ["resumen", "abstract", "summary"]

        for para in self.doc.paragraphs:
            para_lower = para.text.lower().strip()

            if any(k in para_lower for k in keywords):
                capturing = True
                continue

            if capturing:
                if para.style.name.startswith("Heading"):
                    break
                text += para.text + "\n"

        return text.strip()[:2000]

    def get_keywords(self):
        for para in self.doc.paragraphs[:50]:
            text = para.text.lower()
            if "palabras clave" in text or "keywords" in text:
                match = re.search(r"[:\-]\s*(.+?)(?:\n|$)", para.text, re.IGNORECASE)
                if match:
                    return match.group(1).strip()
        return ""

    def get_chapters(self):
        chapters = []
        current_chapter = None
        current_content = []

        for para in self.doc.paragraphs:
            text = para.text.strip()

            if not text:
                continue

            if para.style.name.startswith("Heading") or para.style.name.startswith(
                "List"
            ):
                if (
                    "capítulo" in text.lower()
                    or "capitulo" in text.lower()
                    or re.match(r"^(I|II|III|IV|V|VI|VII|VIII|IX|X|\d+)\.?", text)
                ):
                    if current_chapter:
                        chapters.append(
                            {
                                "titulo": current_chapter,
                                "contenido": "\n".join(current_content),
                            }
                        )

                    current_chapter = text[:300]
                    current_content = []
                else:
                    current_content.append(text)
            else:
                current_content.append(text)

        if current_chapter:
            chapters.append(
                {"titulo": current_chapter, "contenido": "\n".join(current_content)}
            )

        return chapters

    def get_references(self):
        references = []
        in_references = False
        current_ref = []

        keywords = ["referencias", "bibliografía", "bibliografia", "literatura citada"]

        for para in self.doc.paragraphs:
            text = para.text.strip()

            if not in_references:
                if any(k in text.lower() for k in keywords):
                    in_references = True
                continue

            if text:
                current_ref.append(text)
            else:
                if current_ref:
                    ref_text = " ".join(current_ref)

                    tipo = "libro"
                    if "http" in ref_text or "www" in ref_text:
                        tipo = "web"
                    elif re.search(r"\d{4}", ref_text):
                        tipo = "articulo"

                    autor_match = re.match(r"^([^0-9\n]+?)[\d\.\,]", ref_text)

                    references.append(
                        {
                            "tipo": tipo,
                            "autor": autor_match.group(1).strip()
                            if autor_match
                            else ref_text[:100],
                            "titulo": ref_text[:300],
                            "anio": int(
                                re.search(r"\b(19|20)\d{2}\b", ref_text).group()
                            )
                            if re.search(r"\b(19|20)\d{2}\b", ref_text)
                            else None,
                            "fuente": "",
                            "url": re.search(r"https?://[^\s\)]+", ref_text).group()
                            if re.search(r"https?://[^\s\)]+", ref_text)
                            else "",
                        }
                    )

                    current_ref = []

        return references

    def get_year(self):
        # Look for year in document (year often appears at the end like "Barinas, Julio 2025")
        for para in self.doc.paragraphs:
            text = para.text.strip()
            # Match pattern "Month Year" like "Julio 2025"
            match = re.search(
                r"(?:Julio|Junio|Marzo|Abril|Mayo|Agosto|Septiembre|Octubre|Noviembre|Diciembre|Enero|Febrero)\s+(20\d{2}|19\d{2})",
                text,
                re.IGNORECASE,
            )
            if match:
                return int(match.group(1))

            # Also try simple 4-digit year that looks reasonable
            match = re.search(r"\b(20\d{2}|19\d{2})\b", text)
            if match:
                year = int(match.group(1))
                if 1990 <= year <= 2030:  # Reasonable year range
                    return year
        return None

    def extract_all(self):
        return {
            "titulo": self.get_title(),
            "autor_principal": self.get_author(),
            "resumen": self.get_abstract(),
            "palabras_clave": self.get_keywords(),
            "anio": self.get_year(),
            "capitulos": self.get_chapters(),
            "referencias": self.get_references(),
        }
