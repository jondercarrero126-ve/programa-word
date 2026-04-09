import pytest
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from word_extractor.extractor import WordExtractor


class TestWordExtractor:
    @pytest.fixture
    def ruta_tesis(self):
        ruta = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "Proyecto_de_investigacion_Jonder_Carrero_Capitulo_III.docx",
        )
        if os.path.exists(ruta):
            return ruta
        pytest.skip("Archivo de tesis no encontrado")

    def test_inicializacion_exitosa(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        assert extractor is not None
        assert extractor.doc is not None

    def test_inicializacion_archivo_inexistente(self):
        with pytest.raises(Exception):
            WordExtractor("archivo_inexistente.docx")

    def test_get_title(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        titulo = extractor.get_title()
        assert titulo is not None
        assert len(titulo) > 0
        assert titulo != "Sin título", "Deberia extraer un titulo del documento"

    def test_get_author(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        autor = extractor.get_author()
        assert autor is not None
        assert len(autor) > 0

    def test_get_abstract(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        abstract = extractor.get_abstract()
        assert abstract is not None
        assert isinstance(abstract, str)

    def test_get_keywords(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        keywords = extractor.get_keywords()
        assert keywords is not None
        assert isinstance(keywords, str)

    def test_get_year(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        year = extractor.get_year()
        assert year is None or isinstance(year, int), "El ano deberia ser None o entero"

    def test_get_chapters(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        chapters = extractor.get_chapters()
        assert chapters is not None
        assert isinstance(chapters, list)

    def test_get_references(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        references = extractor.get_references()
        assert references is not None
        assert isinstance(references, list)
        assert len(references) > 0, "Deberia encontrar referencias en el documento"

    def test_extract_all(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        datos = extractor.extract_all()

        assert datos is not None
        assert isinstance(datos, dict)
        assert "titulo" in datos
        assert "autor_principal" in datos
        assert "resumen" in datos
        assert "capitulos" in datos
        assert "referencias" in datos

    def test_capitulos_tienen_formato_correcto(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        chapters = extractor.get_chapters()

        for cap in chapters:
            assert "titulo" in cap
            assert "contenido" in cap
            assert isinstance(cap["titulo"], str)
            assert isinstance(cap["contenido"], str)

    def test_referencias_tienen_formato_correcto(self, ruta_tesis):
        extractor = WordExtractor(ruta_tesis)
        references = extractor.get_references()

        for ref in references:
            assert "tipo" in ref
            assert "autor" in ref
            assert "titulo" in ref
            assert ref["tipo"] in ["libro", "articulo", "web", "tesis"]


class TestWordExtractorEdgeCases:
    def test_documento_vacio(self, tmp_path):
        from docx import Document

        doc_path = tmp_path / "vacio.docx"
        doc = Document()
        doc.save(str(doc_path))

        extractor = WordExtractor(str(doc_path))
        titulo = extractor.get_title()
        assert titulo == "Sin título"

    def test_documento_solo_texto(self, tmp_path):
        from docx import Document

        doc_path = tmp_path / "simple.docx"
        doc = Document()
        doc.add_paragraph("Hola mundo")
        doc.core_properties.title = "Documento de Prueba"
        doc.save(str(doc_path))

        extractor = WordExtractor(str(doc_path))
        titulo = extractor.get_title()
        assert titulo == "Documento de Prueba"
